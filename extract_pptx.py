from pptx import Presentation
import json

# Load the presentation
prs = Presentation('/vercel/sandbox/uploads/Lecture 1.pptx')

# Extract all text content
content = []
for slide_num, slide in enumerate(prs.slides, 1):
    slide_content = {
        'slide_number': slide_num,
        'text': []
    }
    
    for shape in slide.shapes:
        if hasattr(shape, "text"):
            if shape.text.strip():
                slide_content['text'].append(shape.text.strip())
    
    if slide_content['text']:
        content.append(slide_content)

# Print the content
for slide in content:
    print(f"\n{'='*80}")
    print(f"SLIDE {slide['slide_number']}")
    print('='*80)
    for text in slide['text']:
        print(text)
        print('-'*40)
